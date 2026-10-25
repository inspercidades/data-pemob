"""Regression checks for the indicator review workbook layout."""

import csv
import importlib.util
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


SCRIPT = Path(__file__).resolve().parents[1] / "review" / "format_indicator_review.py"
SPEC = importlib.util.spec_from_file_location("indicator_formatter", SCRIPT)
formatter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(formatter)


def write_csv(directory, name, rows):
    with (directory / name).open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=rows[0])
        writer.writeheader()
        writer.writerows(rows)


class FormatterTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.path = Path(self.temporary.name)
        write_csv(self.path, "indicator_review.csv", [{
            "code": "1.1.1", "short_question": "Ônibus", "full_question": "Há ônibus?",
            "subject": "Oferta", "type": "sim/não", "coverage_cities": "12",
            "coverage_years": "2019-2025", "response_rate_mean": "0.5",
            "panel_cities": "4", "quality_flags": "", "dashboard_pemob": "sim",
            "tabela_pemob": "não", "include_table": "", "include_dashboard": "",
            "notes": "",
        }])
        write_csv(self.path, "indicator_review_metropolitan.csv", [{
            "code": "1.1.1", "full_question": "Há serviço metropolitano?",
            "coverage_places": "8", "coverage_years": "2019, 2021-2025",
            "response_rate_mean": "0.4", "panel_places": "3", "quality_flags": "",
            "include_table": "", "include_dashboard": "", "notes": "",
        }])
        write_csv(self.path, "indicator_diagnostics.csv", [
            {"level": "municipal", "code": "1.1.1", "n_answers": "20",
             "cities_2019": "12", "places_2019": ""},
            {"level": "metropolitan", "code": "1.1.1", "n_answers": "10",
             "cities_2019": "", "places_2019": "8"},
        ])
        write_csv(self.path, "participants.csv", [{"year": "2025", "participants": "12"}])
        write_csv(self.path, "participants_metropolitan.csv", [{"year": "2025", "participants": "8"}])
        write_csv(self.path, "panel_bounds.csv", [{
            "participated_4_plus": "5", "participated_2021_2025": "2",
        }])
        write_csv(self.path, "panel_bounds_metropolitan.csv", [{
            "participated_4_plus": "4", "participated_2021_2025": "1",
        }])
        write_csv(self.path, "snapshot_long.csv", [{
            "code_muni": "3550308", "name_city": "Cidade", "abbrev_state": "SP",
            "year": "2025",
            "code_indicator": "1.1.1", "value": "SIM", "is_number": "FALSE",
        }])
        write_csv(self.path, "snapshot_long_metropolitan.csv", [{
            "place_name": "Região", "abbrev_state": "SP",
            "year": "2025", "code_indicator": "1.1.1", "value": "SIM",
            "is_number": "FALSE",
        }])
        write_csv(self.path, "snapshot_codes.csv", [{"code_indicator": "1.1.1"}])
        write_csv(self.path, "snapshot_codes_metropolitan.csv", [{"code_indicator": "1.1.1"}])
        write_csv(self.path, "column_labels.csv", [
            {"from": "code", "to": "Código", "definition": "Código da pergunta."},
            {"from": "full_question", "to": "Pergunta completa", "definition": "Texto da pergunta."},
            {"from": "n_answers", "to": "Respostas válidas", "definition": "Total de respostas."},
        ] + [
            {"from": name, "to": label, "definition": f"Definição de {label}."}
            for name, label in {
                "short_question": "Pergunta resumida", "subject": "Tema",
                "type": "Tipo de resposta", "coverage_cities": "Cidades com resposta",
                "coverage_years": "Anos com resposta",
                "response_rate_mean": "Taxa de resposta média",
                "panel_cities": "Cidades com 4+ edições",
                "quality_flags": "Alertas de qualidade",
                "include_table": "Tabela?", "include_dashboard": "Dashboard?",
                "notes": "Notas",
            }.items()
        ])

    def test_separate_compact_review_and_diagnostics(self):
        workbook = formatter.build_workbook(self.path)
        self.assertIn("municipal", workbook.sheetnames)
        self.assertIn("metropolitana", workbook.sheetnames)
        self.assertIn("diagnósticos", workbook.sheetnames)
        municipal = workbook["municipal"]
        metropolitan = workbook["metropolitana"]
        self.assertEqual(municipal["A2"].value, "1.1.1")
        self.assertEqual(metropolitan["A2"].value, "1.1.1")
        self.assertLess(municipal.max_column, 20)
        self.assertLess(metropolitan.max_column, 15)
        municipal_headers = [cell.value for cell in municipal[1]]
        self.assertEqual(municipal_headers, [
            "Código", "Pergunta resumida", "Pergunta completa", "Tema",
            "Tipo de resposta", "Cidades com resposta", "Anos com resposta",
            "Taxa de resposta média", "Alertas de qualidade", "Hoje no site",
            "Tabela?", "Dashboard?", "Notas",
        ])
        self.assertEqual([cell.value for cell in metropolitan[1]], [
            "Código", "Pergunta completa", "Regiões com resposta",
            "Anos com resposta", "Taxa de resposta média",
            "Alertas de qualidade", "Tabela?", "Dashboard?", "Notas",
        ])
        self.assertEqual(
            municipal.cell(2, municipal_headers.index("Hoje no site") + 1).value,
            "Só dashboard",
        )
        for sheet in (municipal, metropolitan):
            headers = [cell.value for cell in sheet[1]]
            validated = {
                str(validation.sqref): validation.formula1
                for validation in sheet.data_validations.dataValidation
            }
            for label in ("Tabela?", "Dashboard?"):
                letter = chr(ord("A") + headers.index(label))
                self.assertEqual(validated[f"{letter}2"], '"sim,não"')
        self.assertNotIn("2020", [cell.value for cell in metropolitan[1]])
        self.assertEqual(workbook["diagnósticos"].max_row, 3)
        diagnostic_headers = [cell.value for cell in workbook["diagnósticos"][1]]
        self.assertEqual(len(diagnostic_headers), len(set(diagnostic_headers)))
        self.assertIn("Municípios 2019", diagnostic_headers)
        self.assertIn("Regiões 2019", diagnostic_headers)
        self.assertEqual(
            [cell.value for cell in workbook["dados municipais"][1]],
            ["Ano", "Código IBGE", "Município", "UF", "1.1.1"],
        )
        self.assertEqual(
            [cell.value for cell in workbook["dados municipais"][2]],
            [2025, 3550308, "Cidade", "SP", "SIM"],
        )
        self.assertEqual(
            [cell.value for cell in workbook["dados metropolitanos"][1]],
            ["Ano", "Região metropolitana", "UF", "1.1.1"],
        )

    def test_preserves_reviewer_entries_by_level_and_code(self):
        old = Workbook()
        sheet = old.active
        sheet.title = "municipal"
        sheet.append(["Código", "Tabela?", "Dashboard?", "Notas"])
        sheet.append(["1.1.1", "sim", "não", "Conferir fonte"])
        old_path = self.path / "old.xlsx"
        old.save(old_path)
        workbook = formatter.build_workbook(self.path, existing_path=old_path)
        sheet = workbook["municipal"]
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet.cell(2, headers.index("Tabela?") + 1).value, "sim")
        self.assertEqual(sheet.cell(2, headers.index("Dashboard?") + 1).value, "não")
        self.assertEqual(sheet.cell(2, headers.index("Notas") + 1).value, "Conferir fonte")
        metro = workbook["metropolitana"]
        metro_headers = [cell.value for cell in metro[1]]
        self.assertIsNone(metro.cell(2, metro_headers.index("Tabela?") + 1).value)

    def legacy_workbook(self, interest, include, notes):
        old = Workbook()
        sheet = old.active
        sheet.title = "municipal"
        sheet.append(["Código", "Interesse", "Incluir?", "Notas"])
        sheet.append(["1.1.1", interest, include, notes])
        old_path = self.path / "legacy.xlsx"
        old.save(old_path)
        return old_path

    def test_migrates_notes_from_legacy_reviewer_columns(self):
        old_path = self.legacy_workbook(None, None, "Conferir fonte")
        workbook = formatter.build_workbook(self.path, existing_path=old_path)
        sheet = workbook["municipal"]
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(sheet.cell(2, headers.index("Notas") + 1).value, "Conferir fonte")
        self.assertIsNone(sheet.cell(2, headers.index("Tabela?") + 1).value)
        self.assertIsNone(sheet.cell(2, headers.index("Dashboard?") + 1).value)

    def test_refuses_to_drop_legacy_decisions(self):
        old_path = self.legacy_workbook("alta", "sim", "Conferir fonte")
        with self.assertRaisesRegex(ValueError, "Legacy reviewer decisions"):
            formatter.build_workbook(self.path, existing_path=old_path)

    def test_fails_if_reviewed_question_disappears(self):
        old = Workbook()
        sheet = old.active
        sheet.title = "revisão"
        sheet.append(["Código", "Tabela?", "Dashboard?", "Notas"])
        sheet.append(["9.9.9", "sim", "sim", "Keep this decision"])
        old_path = self.path / "old.xlsx"
        old.save(old_path)
        with self.assertRaisesRegex(ValueError, "Reviewed keys missing"):
            formatter.build_workbook(self.path, existing_path=old_path)

    def test_fails_on_duplicate_level_code(self):
        rows = formatter.read_csv(self.path, "indicator_review_metropolitan.csv")
        write_csv(self.path, "indicator_review_metropolitan.csv", rows + rows)
        with self.assertRaisesRegex(ValueError, "Duplicate review key"):
            formatter.build_workbook(self.path)


if __name__ == "__main__":
    unittest.main()
